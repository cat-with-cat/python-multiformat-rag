from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def _pdf_literal(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")


def create_text_pdf(path: Path, page_texts: list[str]) -> None:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)

    for text in page_texts:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(_pdf_literal(text))
        page[NameObject("/Contents")] = writer._add_object(stream)

    with path.open("wb") as file_handle:
        writer.write(file_handle)


def create_presentation(path: Path, slide_texts: list[str]) -> None:
    presentation = Presentation()
    for text in slide_texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        text_box.text = text
    presentation.save(path)


@pytest.fixture
def synthetic_documents(tmp_path: Path) -> Path:
    data_dir = tmp_path / "documents"
    data_dir.mkdir()

    (data_dir / "guide.txt").write_text(
        "Knowledge Base\n\nThe system supports traceable retrieval.",
        encoding="utf-8",
    )
    (data_dir / "products.csv").write_text(
        "name,feature\nAlpha,semantic search\nBeta,source citations\n",
        encoding="utf-8",
    )
    create_text_pdf(
        data_dir / "manual.pdf",
        ["PDF page one describes indexing.", "PDF page two describes persistence."],
    )
    create_presentation(
        data_dir / "slides.pptx",
        ["Slide one: document parsing", "Slide two: grounded answers"],
    )
    return data_dir

