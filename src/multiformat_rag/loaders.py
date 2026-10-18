"""Deterministic loaders for TXT, CSV, PDF, and PPTX documents."""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path

from llama_index.core import Document
from pptx import Presentation
from pypdf import PdfReader

from multiformat_rag.exceptions import DocumentLoadError, UnsupportedFormatError

SUPPORTED_EXTENSIONS = frozenset({".txt", ".csv", ".pdf", ".pptx"})


def _source_id(source_file: str, locator: str) -> str:
    value = f"{source_file}:{locator}".encode("utf-8")
    return hashlib.sha256(value).hexdigest()[:20]


class MultiFormatDocumentLoader:
    """Load supported files into a common LlamaIndex ``Document`` model.

    Paths stored in metadata are relative to the requested input root so a
    persisted index remains portable and does not expose a developer's home
    directory.
    """

    def load(self, input_path: str | Path) -> list[Document]:
        path = Path(input_path).expanduser()
        if not path.exists():
            raise DocumentLoadError(f"输入路径不存在：{path}")

        if path.is_file():
            return self._load_file(path, source_file=path.name)

        files = sorted(
            candidate
            for candidate in path.rglob("*")
            if candidate.is_file() and candidate.suffix.lower() in SUPPORTED_EXTENSIONS
        )
        if not files:
            supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise DocumentLoadError(f"目录中没有受支持的文件（{supported}）：{path}")

        documents: list[Document] = []
        for file_path in files:
            source_file = file_path.relative_to(path).as_posix()
            documents.extend(self._load_file(file_path, source_file=source_file))
        return documents

    def _load_file(self, path: Path, source_file: str) -> list[Document]:
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFormatError(
                f"不支持的文档格式：{extension or '<无扩展名>'}"
            )

        try:
            if extension == ".txt":
                return self._load_txt(path, source_file)
            if extension == ".csv":
                return self._load_csv(path, source_file)
            if extension == ".pdf":
                return self._load_pdf(path, source_file)
            return self._load_pptx(path, source_file)
        except DocumentLoadError:
            raise
        except Exception as exc:
            raise DocumentLoadError(f"解析文档失败：{source_file}（{exc}）") from exc

    @staticmethod
    def _base_metadata(source_file: str, file_type: str, locator: str) -> dict:
        return {
            "source_file": source_file,
            "file_type": file_type,
            "source_id": _source_id(source_file, locator),
        }

    def _load_txt(self, path: Path, source_file: str) -> list[Document]:
        text = path.read_text(encoding="utf-8-sig")
        metadata = self._base_metadata(source_file, "txt", "document")
        return [Document(text=text, metadata=metadata)]

    def _load_csv(self, path: Path, source_file: str) -> list[Document]:
        documents: list[Document] = []
        with path.open("r", encoding="utf-8-sig", newline="") as file_handle:
            reader = csv.DictReader(file_handle)
            if not reader.fieldnames:
                return []
            for row_number, row in enumerate(reader, start=2):
                fields = [
                    f"{column.strip()}: {(value or '').strip()}"
                    for column, value in row.items()
                    if column is not None and (value or "").strip()
                ]
                metadata = self._base_metadata(
                    source_file, "csv", f"row:{row_number}"
                )
                metadata["row_number"] = row_number
                documents.append(Document(text="\n".join(fields), metadata=metadata))
        return documents

    def _load_pdf(self, path: Path, source_file: str) -> list[Document]:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise DocumentLoadError(f"PDF 已加密且无法读取：{source_file}") from exc

        documents: list[Document] = []
        for page_number, page in enumerate(reader.pages, start=1):
            metadata = self._base_metadata(
                source_file, "pdf", f"page:{page_number}"
            )
            metadata["page_number"] = page_number
            documents.append(
                Document(text=page.extract_text() or "", metadata=metadata)
            )
        return documents

    def _load_pptx(self, path: Path, source_file: str) -> list[Document]:
        presentation = Presentation(str(path))
        documents: list[Document] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        parts.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values = [cell.text.strip() for cell in row.cells]
                        if any(values):
                            parts.append(" | ".join(values))

            metadata = self._base_metadata(
                source_file, "pptx", f"slide:{slide_number}"
            )
            metadata["slide_number"] = slide_number
            documents.append(Document(text="\n".join(parts), metadata=metadata))
        return documents
